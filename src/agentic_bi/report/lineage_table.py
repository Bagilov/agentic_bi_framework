"""Excel view of the flat lineage matrix.

Consumes the rows produced by :func:`agentic_bi.report.lineage_matrix.lineage_to_table`
and renders an xlsx workbook with three sheets:

    Lineage  - the full matrix (autofilter, per-level cell palette)
    Summary  - node counters per level/type + brand-book summary
    Columns  - static documentation of the matrix format

Requires ``pandas`` and ``openpyxl`` (the ``report`` optional extra).
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from agentic_bi.report.lineage_matrix import COLUMNS, lineage_to_table

# Light fill per level, used for the Excel line palette.
LEVEL_COLORS = {
    "L1": "D5E8D4",
    "L2": "FFF9C4",
    "L3": "C8E6C9",
    "L4": "B3E5FC",
    "L5": "D1C4E9",
    "L6": "FFE0B2",
    "L7": "F8BBD0",
}

_COLUMN_WIDTH = {
    "Level": 6,
    "Layer": 18,
    "L1 · Origin": 30,
    "L2 · Source": 20,
    "L3 · Power Query": 22,
    "L4 · Table": 20,
    "L5 · Semantics": 22,
    "L6 · Page": 26,
    "L7 · Visual": 30,
    "Parent": 26,
    "Type": 20,
    "Detail / Source": 40,
    "Refs": 14,
    "Came from": 26,
    "Goes to": 26,
    "Brand · fonts, %": 18,
    "Brand · colors, %": 17,
    "Brand · status": 18,
    "Brand · notes": 36,
}


class LineageTable:
    """Flat lineage matrix from a ``PbiLineage.to_json()`` dict.

    Args:
        data: dict in the ``PbiLineage.to_json()`` structure.
    """

    def __init__(self, data: dict):
        self.data = data
        self.rows = lineage_to_table(data)

    @property
    def dataframe(self):
        """The matrix as a pandas DataFrame (lazy import)."""
        import pandas as pd

        return pd.DataFrame(self.rows, columns=COLUMNS).fillna("")

    @property
    def summary(self) -> dict:
        """Node counters per level/type plus the brand-book summary."""
        df = self.dataframe
        return {
            "nodes": len(df),
            "by_level": df.groupby("Layer").size().to_dict(),
            "by_type": df["Type"].value_counts().to_dict(),
            "with_refs": int((df["Refs"] > 0).sum()),
            "orphans": [
                r for r in self.rows if not r.get("Came from") and r.get("Level", 0) > 1
            ],
            "brand": (self.data.get("brand") or {}).get("summary", {}),
        }

    def to_excel(self, out_path: str | Path) -> Path:
        """Save the matrix to xlsx (Lineage + Summary + Columns)."""
        import pandas as pd
        from openpyxl.styles import Font, PatternFill

        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        df = self.dataframe

        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Lineage", index=False)
            wb = writer.book
            ws = wb["Lineage"]

            for cell in ws[1]:
                cell.fill = PatternFill(
                    start_color="2C3E50", end_color="2C3E50", fill_type="solid"
                )
                cell.font = Font(color="FFFFFF", bold=True)
            for col in ws.columns:
                header = col[0].value
                ws.column_dimensions[col[0].column_letter].width = _COLUMN_WIDTH.get(
                    header, 14
                )
            ws.auto_filter.ref = ws.dimensions

            header_to_col = {h: i for i, h in enumerate(COLUMNS, start=1)}
            fill_by_col = {
                "L1 · Origin": "L1",
                "L2 · Source": "L2",
                "L3 · Power Query": "L3",
                "L4 · Table": "L4",
                "L5 · Semantics": "L5",
                "L6 · Page": "L6",
                "L7 · Visual": "L7",
            }
            for row in ws.iter_rows(min_row=2):
                for header, key in fill_by_col.items():
                    cell = row[header_to_col[header] - 1]
                    if cell.value:
                        cell.fill = PatternFill(
                            start_color=LEVEL_COLORS[key],
                            end_color=LEVEL_COLORS[key],
                            fill_type="solid",
                        )

            self._write_summary_sheet(wb, self.summary)
            self._write_columns_sheet(wb)
        return out

    @staticmethod
    def _write_columns_sheet(wb) -> None:
        """Document the matrix format (Column / Description / Source)."""
        rows = [
            ["Lineage matrix format"],
            [],
            ["Column", "Description", "Source"],
            ["Level", "Lineage layer number (1-7)", "computed"],
            [
                "Layer",
                "Layer name (Origin/Source/Power Query/Table/Semantics/Page/Visual)",
                "LEVEL",
            ],
            ["L1..L7", "Node of the matching layer (one per row)", "lineage JSON"],
            ["Parent", "Node one layer above", "layer links"],
            ["Type", "Node type (origin category/source/visual type)", "lineage JSON"],
            [
                "Detail / Source",
                "Path/DSN/columns/measures or bindings",
                "lineage JSON",
            ],
            ["Refs", "How many times the node is used in visuals", "lineage JSON refs"],
            ["Came from", "Input node (layer N-1)", "layer links"],
            ["Goes to", "Output node (layer N+1)", "layer links"],
            [
                "Brand · fonts, %",
                "Brand-font compliance of the page/visual",
                "brand block",
            ],
            ["Brand · colors, %", "Brand-palette color compliance", "brand block"],
            [
                "Brand · status",
                "compliant / warnings / violation / no data",
                "brand block",
            ],
            [
                "Brand · notes",
                "Deviations: off-brand fonts/colors, missing brand red",
                "brand block",
            ],
            [],
            [
                "Note: level 7 rows carry element metrics, level 6 page metrics, "
                "level 1 the dashboard summary."
            ],
        ]
        ws = wb.create_sheet("Columns")
        for row in rows:
            ws.append(row)

    @staticmethod
    def _write_summary_sheet(wb, summ: dict) -> None:
        s = wb.create_sheet("Summary")
        s.append(["Lineage matrix statistics"])
        s.append([])
        s.append(["Total nodes", summ["nodes"]])
        s.append(["Nodes with refs>0", summ["with_refs"]])
        s.append([])
        s.append(["By level"])
        for lvl, cnt in summ["by_level"].items():
            s.append([lvl, cnt])
        s.append([])
        s.append(["By type"])
        for t, cnt in sorted(summ["by_type"].items(), key=lambda kv: -kv[1]):
            s.append([t, cnt])
        if summ["orphans"]:
            s.append([])
            s.append(["Orphan nodes (no source)"])
            for o in summ["orphans"]:
                s.append([o.get("Level"), o.get("Layer"), o.get("Detail / Source")])
        b = summ.get("brand") or {}
        if b:
            s.append([])
            s.append(["Brand-book compliance (page elements)"])
            s.append(["Pages", b.get("pages", 0)])
            s.append(["Compliant", b.get("ok", 0)])
            s.append(["Warnings", b.get("warnings", 0)])
            s.append(["Violations", b.get("errors", 0)])
            s.append(["No data", b.get("unknown", 0)])
            s.append(["Fonts, %", b.get("fonts_compliance")])
            s.append(["Colors, %", b.get("colors_compliance")])
            if b.get("error_pages"):
                s.append(["Pages with violations", ", ".join(b["error_pages"])])


def load_json(path: str | Path) -> dict:
    """Load a lineage JSON file, or generate it from a ``.pbix`` path."""
    path = Path(path)
    if path.suffix.lower() == ".pbix":
        from agentic_bi.lineage.pbi_lineage import PbiLineage

        with PbiLineage(path) as lg:
            return lg.to_json()
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Flat level matrix of data lineage (JSON/.pbix -> xlsx)."
    )
    parser.add_argument("input", help="path to lineage.json or .pbix")
    parser.add_argument(
        "-o",
        "--output",
        default=None,
        help="output xlsx (default: '<name>_table.xlsx' next to input)",
    )
    args = parser.parse_args()

    data = load_json(args.input)
    inp = Path(args.input)
    out = Path(args.output) if args.output else inp.with_name(inp.stem + "_table.xlsx")

    table = LineageTable(data)
    table.to_excel(out)
    print(f"Lineage matrix saved: {out}")
    print(f"Total nodes: {len(table.rows)}")


if __name__ == "__main__":
    main()
